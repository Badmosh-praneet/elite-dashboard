import openpyxl
import warnings

def modify_dsr():
    input_path = 'Sample DSR.xlsx'
    output_path = 'Test DSR.xlsx'
    print(f'Loading {input_path}...')
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(input_path)
    
    ws = wb['Leads']
    
    # Replace Lokesh Reddy K with Shivani Reddy K
    replaced_count = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and 'Lokesh Reddy K' in cell.value:
                cell.value = cell.value.replace('Lokesh Reddy K', 'Shivani Reddy K')
                replaced_count += 1
    
    print(f'Replaced {replaced_count} occurrences of Lokesh Reddy K')
    
    # Copy the last row and add new leads
    max_row = ws.max_row
    last_row_data = [cell.value for cell in ws[max_row]]
    
    new_names = ['Shivani Reddy K', 'Test Lead Beta', 'Test Lead Gamma']
    for i, name in enumerate(new_names):
        new_row_data = list(last_row_data)
        # Find Name column
        name_col_idx = None
        for j, cell in enumerate(ws[1]):
            if cell.value and 'Name' in str(cell.value):
                name_col_idx = j
                break
        
        if name_col_idx is not None:
            new_row_data[name_col_idx] = name
        
        ws.append(new_row_data)
        
    print(f'Added {len(new_names)} new leads')
    
    print(f'Saving to {output_path}...')
    wb.save(output_path)
    print('Done!')

if __name__ == '__main__':
    modify_dsr()

