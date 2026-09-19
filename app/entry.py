"""
Data-entry and live-stream routes.

Split out from main.py so the read API stays easy to scan: everything that
changes the database, plus the event stream the dashboard listens on, lives here
and is mounted onto the app as a router.

Supports dual-mode execution: writes directly to PostgreSQL when configured,
and maintains in-memory and client-side reactive state when running in serverless/offline mode.
"""

from __future__ import annotations

import calendar
import io
import json
import logging
import os
import re
from datetime import date, datetime
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
import openpyxl
import psycopg
from psycopg.errors import IntegrityError

from . import fallback
from .db import (is_db_ready, fetch_all, pool, DSN, ensure_pool_open,
                 connect as db_connect, session)
from .events import broker
from etl.dimensions import activate_period
from etl.load_dsr import Loader
from .write import (AllotmentIn, BookingIn, BookingPatch, LeadIn, RegistrationIn,
                    TestDriveIn, VehicleIn, create_allotment, create_booking,
                    create_lead, create_registration, create_test_drive,
                    delete_row, update_booking, upsert_vehicle)

log = logging.getLogger("dsr.entry")
router = APIRouter()

# Everything a reporting month owns. Facts are keyed by the load that produced
# them (load_period_id); lead and booking additionally carry a business
# period_id, which is where hand-entered rows are attached.
PERIOD_FACTS = ["registration", "allotment", "booking", "test_drive", "lead"]
PERIOD_TARGET_TABLES = [
    "target_daily_tracker", "target_booking_commitment",
    "target_channel_funnel", "target_consultant_scorecard",
]

MONTH_MAP = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "september": 9, "oct": 10, "october": 10,
    "nov": 11, "november": 11, "dec": 12, "december": 12,
}


def _resolve_period(period: str | None, filename: str) -> tuple[str, date, date]:
    """
    Work out which month a report is for - or refuse to guess.

    This used to default to August 2026 whenever it could not tell. A file
    called "Test DSR.xlsx" therefore loaded silently into AUG2026 and replaced
    that month's bookings, test drives, allotments and registrations. Getting
    this wrong destroys a month of data, so an unanswerable case is now an
    error the uploader can act on rather than a guess nobody sees.
    """
    month = year = None

    # An explicit label is the most reliable signal, and it is the one form the
    # scan below cannot read: in "OCT2026" the year is glued to the month, so
    # there is no word boundary for \b(20\d\d)\b to find.
    tag = re.match(r"^\s*([a-z]{3,9})\s*[-_ ]?\s*(20\d\d)\s*$", (period or "").lower())
    if tag and tag.group(1) in MONTH_MAP:
        month, year = MONTH_MAP[tag.group(1)], int(tag.group(2))
    else:
        text = f"{period or ''} {filename}".lower()
        for name, num in MONTH_MAP.items():
            if re.search(r"\b" + name, text):
                month = num
                break
        ym = re.search(r"\b(20\d\d)\b", text)
        year = int(ym.group(1)) if ym else None

    if month is None or year is None:
        missing = "month" if month is None else "year"
        raise HTTPException(
            400,
            f"Could not tell which {missing} this report covers, and guessing would "
            f"overwrite whichever month it guessed. Set the period explicitly "
            f"(for example OCT2026), or name the file with its month and year "
            f"(for example 'DSR October 2026.xlsx').",
        )

    abbr = calendar.month_abbr[month].upper()
    # Always the canonical ABBRYYYY. A label typed "OCT 2026" would otherwise
    # create a second period alongside "OCT2026", each holding half the month.
    label = f"{abbr}{year}"

    _, last_day = calendar.monthrange(year, month)
    return label, date(year, month, 1), date(year, month, last_day)


# =====================================================================
# Live change stream
# =====================================================================

@router.get("/api/events", tags=["live"], include_in_schema=False)
async def events():
    """
    Server-sent events: one `change` frame whenever the database moves.
    """
    if os.environ.get("VERCEL") or not is_db_ready():
        async def vercel_stream():
            yield 'event: ready\ndata: {"connected": false, "serverless": true}\n\n'
        return StreamingResponse(
            vercel_stream(),
            media_type="text/event-stream",
            headers={"Cache-Control": "no-cache"},
        )

    return StreamingResponse(
        broker.stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/api/live-status", tags=["live"])
def live_status():
    """Whether the listener is attached, and how many dashboards are watching."""
    if os.environ.get("VERCEL") or not is_db_ready():
        return {"connected": False, "serverless": True, "subscribers": 0, "events_seen": 0, "last_error": None}
    return broker.status


# =====================================================================
# Writes
# =====================================================================

def _commit(fn, *args):
    """Run one writer inside a transaction and turn its errors into HTTP codes."""
    ensure_pool_open()
    try:
        with pool.connection(timeout=10.0) as cx:
            with cx.transaction():
                return fn(cx, *args)
    except HTTPException:
        raise
    except PermissionError as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    except IntegrityError as exc:
        raise HTTPException(409, str(exc).strip().splitlines()[0]) from exc
    except Exception as exc:
        log.exception("Database transaction failed: %s", exc)
        raise HTTPException(500, f"Database transaction failed: {exc}") from exc


@router.post("/api/leads", status_code=201, tags=["entry"])
def add_lead(body: LeadIn):
    if is_db_ready():
        res = _commit(create_lead, body)
        broker.notify_sync("lead")
        return res
    return fallback.store.add_lead(body.model_dump())


@router.post("/api/bookings", status_code=201, tags=["entry"])
def add_booking(body: BookingIn):
    if is_db_ready():
        res = _commit(create_booking, body)
        broker.notify_sync("booking")
        return res
    return fallback.store.add_booking(body.model_dump())


@router.patch("/api/bookings/{booking_id}", tags=["entry"])
def patch_booking(booking_id: int, body: BookingPatch):
    if is_db_ready():
        res = _commit(update_booking, booking_id, body)
        broker.notify_sync("booking")
        return res
    return fallback.store.patch_booking(booking_id, body.model_dump(exclude_unset=True))


@router.post("/api/test-drives", status_code=201, tags=["entry"])
def add_test_drive(body: TestDriveIn):
    if is_db_ready():
        res = _commit(create_test_drive, body)
        broker.notify_sync("test_drive")
        return res
    return fallback.store.add_test_drive(body.model_dump())


@router.post("/api/vehicles", status_code=201, tags=["entry"])
def add_vehicle(body: VehicleIn):
    if is_db_ready():
        res = _commit(upsert_vehicle, body)
        broker.notify_sync("vehicle")
        return res
    return fallback.store.add_vehicle(body.model_dump())


@router.post("/api/allotments", status_code=201, tags=["entry"])
def add_allotment(body: AllotmentIn):
    if is_db_ready():
        res = _commit(create_allotment, body)
        broker.notify_sync("allotment")
        return res
    return fallback.store.add_allotment(body.model_dump())


@router.post("/api/registrations", status_code=201, tags=["entry"])
def add_registration(body: RegistrationIn):
    if is_db_ready():
        res = _commit(create_registration, body)
        broker.notify_sync("registration")
        return res
    return fallback.store.add_registration(body.model_dump())


@router.delete("/api/entries/{table}/{row_id}", status_code=200, tags=["entry"])
def delete_entry(table: str, row_id: int):
    if is_db_ready():
        res = _commit(delete_row, table, row_id)
        broker.notify_sync(table)
        return res
    return fallback.store.delete_entry(table, row_id)


# =====================================================================
# Dropdowns and datalists for the entry drawer
# =====================================================================

@router.get("/api/entry-options", tags=["entry"])
def entry_options():
    if is_db_ready():
        try:
            # Seven queries over one checked-out connection rather than seven:
            # against a remote database the checkout costs more than the query.
            with session() as cx:
                q = lambda sql: cx.execute(sql).fetchall()
                return {
                    "consultants": [r["display_name"] for r in q(
                        "SELECT display_name FROM dim_consultant WHERE is_active "
                        "ORDER BY display_name")],
                    "sources": [r["name"] for r in q(
                        "SELECT name FROM dim_lead_source ORDER BY name")],
                    "models": [r["name"] for r in q(
                        "SELECT name FROM dim_model ORDER BY name")],
                    "variants": q(
                        "SELECT m.name AS model, v.name AS variant, v.long_model_text "
                        "FROM dim_variant v JOIN dim_model m USING (model_id) "
                        "ORDER BY m.name, v.name"),
                    "colours": [r["name"] for r in q(
                        "SELECT name FROM dim_colour ORDER BY name")],
                    "fulfilment_statuses": ["BOOKED", "NO_STOCK", "ALLOTED",
                                            "RETAILED", "CANCELLED"],
                    "open_bookings": q("""
                        SELECT b.booking_id, b.customer_name, m.name AS model,
                               dv.name AS variant, c.display_name AS consultant
                        FROM booking b
                        LEFT JOIN dim_model m      ON m.model_id = b.model_id
                        LEFT JOIN dim_variant dv   ON dv.variant_id = b.variant_id
                        LEFT JOIN dim_consultant c ON c.consultant_id = b.consultant_id
                        WHERE b.fulfilment_status IN ('BOOKED', 'NO_STOCK')
                        ORDER BY b.booking_date DESC NULLS LAST
                        LIMIT 200
                    """),
                    "free_chassis": q("""
                        SELECT chassis_number, model, variant, colour, stock_aging_days
                        FROM v_stock WHERE stock_status = 'FREESTOCK'
                        ORDER BY stock_aging_days DESC NULLS LAST
                    """),
                }
        except Exception:
            pass
    return fallback.get_entry_options()


@router.get("/api/recent-activity", tags=["entry"])
def recent_activity(limit: int = Query(25, le=100)):
    if is_db_ready():
        try:
            return fetch_all("""
                SELECT 'booking' AS kind, booking_id AS id, customer_name AS who,
                       loaded_at, entered_by
                  FROM booking WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'lead', lead_id, lead_name, loaded_at, entered_by
                  FROM lead WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'test drive', test_drive_id, lead_name, loaded_at, entered_by
                  FROM test_drive WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'allotment', allotment_id, customer_name, loaded_at, entered_by
                  FROM allotment WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'registration', registration_id, customer_name, loaded_at, entered_by
                  FROM registration WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'vehicle', vehicle_id, chassis_number, loaded_at, entered_by
                  FROM vehicle WHERE origin = 'MANUAL'
                ORDER BY loaded_at DESC
                LIMIT %s
            """, (limit,))
        except Exception:
            pass
    return fallback.store.get_recent_activity(limit)


# =====================================================================
# Reporting period
# =====================================================================

@router.get("/api/periods", tags=["entry"])
def periods():
    if is_db_ready():
        try:
            res = fetch_all("""
                SELECT p.label, p.period_start, p.period_end, p.is_active,
                       (SELECT count(*) FROM lead l    WHERE l.period_id = p.period_id) AS leads,
                       (SELECT count(*) FROM booking b WHERE b.period_id = p.period_id) AS bookings
                FROM dim_period p
                ORDER BY p.period_start DESC
            """)
            if res:
                return res
        except Exception as e:
            import traceback
            traceback.print_exc()
            raise
    return fallback.get_periods()


@router.post("/api/period/{label}/activate", tags=["entry"])
def set_active_period(label: str):
    if is_db_ready():
        try:
            return _commit(activate_period, label)
        except Exception as e:
            import traceback
            traceback.print_exc()
            raise HTTPException(status_code=500, detail=str(e))
    return {"activated": label}


@router.get("/api/periods/{label}/contents", tags=["entry"])
def period_contents(label: str):
    """
    What deleting this month would remove, so the confirmation can say so
    rather than asking the user to take it on faith.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")
    with session() as cx:
        row = cx.execute(
            "SELECT period_id, label, is_active FROM dim_period WHERE upper(label) = upper(%s)",
            (label,)).fetchone()
        if not row:
            raise HTTPException(404, f"There is no month called {label}.")
        pid = row["period_id"]
        counts = {}
        for table in PERIOD_FACTS:
            counts[table] = cx.execute(
                f"SELECT count(*) AS n FROM {table} "
                f"WHERE load_period_id = %s OR period_id = %s"
                if table in ("lead", "booking") else
                f"SELECT count(*) AS n FROM {table} WHERE load_period_id = %s",
                (pid, pid) if table in ("lead", "booking") else (pid,)).fetchone()["n"]
        manual = cx.execute(
            "SELECT count(*) AS n FROM lead WHERE origin = 'MANUAL' AND period_id = %s",
            (pid,)).fetchone()["n"] + cx.execute(
            "SELECT count(*) AS n FROM booking WHERE origin = 'MANUAL' AND period_id = %s",
            (pid,)).fetchone()["n"]
        return {
            "label": row["label"],
            "is_active": row["is_active"],
            "counts": {k: v for k, v in counts.items() if v},
            "total": sum(counts.values()),
            "hand_entered": manual,
        }


@router.post("/api/periods/{label}/clear", tags=["entry"])
def clear_period(label: str, confirm: str = Query(..., description="Must repeat the month's label")):
    """
    Empty a month without removing it.

    Distinct from deleting the month outright: the month, its date range and its
    place in the switcher survive, so the dashboard keeps working and shows
    zeros until a workbook is uploaded for it. This is the "start this month
    again" button, where delete is "this month should not exist".

    Uploading a workbook already replaces the month it belongs to, so this is
    only needed when someone wants the month emptied WITHOUT loading anything
    in its place.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")
    if confirm.strip().upper() != label.strip().upper():
        raise HTTPException(400, "Type the month's label exactly to confirm.")

    removed: dict[str, int] = {}
    with db_connect() as cx:
        with cx.transaction():
            row = cx.execute(
                "SELECT period_id, label FROM dim_period WHERE upper(label) = upper(%s) FOR UPDATE",
                (label,)).fetchone()
            if not row:
                raise HTTPException(404, f"There is no month called {label}.")
            pid, real_label = row[0], row[1]

            for table in PERIOD_FACTS:
                if table in ("lead", "booking"):
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s OR period_id = %s",
                        (pid, pid)).rowcount
                else:
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            for table in PERIOD_TARGET_TABLES:
                n = cx.execute(f"DELETE FROM {table} WHERE period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n

    broker.notify_sync("lead")
    broker.notify_sync("booking")
    return {"cleared": real_label, "removed": removed, "total": sum(removed.values())}


@router.delete("/api/periods/{label}", tags=["entry"])
def delete_period(label: str, confirm: str = Query(..., description="Must repeat the month's label")):
    """
    Remove a reporting month and everything loaded under it.

    This cannot be undone from the dashboard - the only way back is to upload
    that month's workbook again - so it is guarded three ways: the caller has to
    repeat the label, the month must not be the active one, and the last
    remaining month cannot be removed.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")

    if confirm.strip().upper() != label.strip().upper():
        raise HTTPException(400, "Type the month's label exactly to confirm the deletion.")

    removed: dict[str, int] = {}
    with db_connect(autocommit=False) as cx:
        with cx.transaction():
            row = cx.execute(
                "SELECT period_id, label, is_active FROM dim_period "
                "WHERE upper(label) = upper(%s) FOR UPDATE",
                (label,)).fetchone()
            if not row:
                raise HTTPException(404, f"There is no month called {label}.")
            pid, real_label, is_active = row[0], row[1], row[2]

            if is_active:
                raise HTTPException(
                    409,
                    f"{real_label} is the month the dashboard is currently reporting on. "
                    f"Switch to another month first, then delete it.")

            if cx.execute("SELECT count(*) FROM dim_period").fetchone()[0] <= 1:
                raise HTTPException(409, "This is the only month on record; it cannot be removed.")

            # Facts first, then the targets, then the month itself, so no foreign
            # key is left pointing at a row that has gone.
            for table in PERIOD_FACTS:
                if table in ("lead", "booking"):
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s OR period_id = %s",
                        (pid, pid)).rowcount
                else:
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            for table in PERIOD_TARGET_TABLES:
                n = cx.execute(f"DELETE FROM {table} WHERE period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            cx.execute("DELETE FROM etl_run WHERE notes LIKE %s", (f"%[{real_label}]%",))
            cx.execute("DELETE FROM dim_period WHERE period_id = %s", (pid,))

    broker.notify_sync("lead")
    broker.notify_sync("booking")
    return {
        "deleted": real_label,
        "removed": removed,
        "total": sum(removed.values()),
    }


# =====================================================================
# Excel Workbook Ingestion Pipeline
# =====================================================================

@router.post("/api/upload-dsr", tags=["ingestion"])
@router.post("/api/upload-report", tags=["ingestion"])
async def upload_dsr_workbook(
    file: UploadFile = File(..., description="DSR Report file (.xlsx, .xlsm, .csv, .txt, .tsv)"),
    period: str | None = Form(None, description="Optional period label e.g. AUG2026, SEP2026"),
    uploaded_by: str = Form("Reporting Agent", description="Name of agent or manager uploading"),
    table_type: str | None = Form(None, description="Optional target table: auto, booking, lead, vehicle"),
):
    """
    Ingest a DSR report file (Excel workbook, CSV, or Text format).
    Rebuilds/updates facts (leads, bookings, vehicles, test drives),
    updates period alignment in Supabase, and notifies connected web clients.
    """
    fname = file.filename or "unknown_report.txt"
    fname_lower = fname.lower()
    allowed_exts = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".tsv")
    if not fname_lower.endswith(allowed_exts):
        raise HTTPException(
            400,
            f"Invalid file format. Please upload an Excel (.xlsx, .xlsm), CSV (.csv), or Text (.txt, .tsv) file."
        )

    content = await file.read()
    if not content:
        raise HTTPException(400, "Uploaded file is empty.")

    period_label, period_start, period_end = _resolve_period(period, fname)

    # 1. Handle Plain Text, CSV, and TSV files
    if fname_lower.endswith((".csv", ".txt", ".tsv")):
        from etl.text_parser import ingest_text_report
        try:
            parse_result = ingest_text_report(
                content=content,
                filename=fname,
                period_label=period_label,
                period_start=period_start,
                period_end=period_end,
                uploaded_by=uploaded_by,
                table_type=table_type,
            )
            counts = parse_result.get("counts", {})
            warnings = parse_result.get("warnings", [])

            if is_db_ready():
                try:
                    with db_connect(autocommit=True) as cx:
                        cx.execute("""
                            INSERT INTO etl_run (source_file, file_modified, finished_at, row_counts, notes)
                            VALUES (%s, now(), now(), %s, %s)
                        """, (
                            fname,
                            json.dumps(counts),
                            f"Uploaded {parse_result.get('detected_format', 'text')} by {uploaded_by}",
                        ))
                except Exception:
                    pass

            try:
                await broker.notify("workbook_reload", "etl_run", {"counts": counts, "period": period_label})
            except Exception:
                pass

            return {
                "status": "success",
                "message": f"Successfully ingested {parse_result.get('detected_format', 'report')} '{fname}' for {period_label}.",
                "filename": fname,
                "file_type": "text/csv",
                "detected_format": parse_result.get("detected_format"),
                "detected_table": parse_result.get("detected_table"),
                "period": period_label,
                "period_range": f"{period_start} to {period_end}",
                "uploaded_by": uploaded_by,
                "counts": counts,
                "warnings": warnings,
                "environment": "postgres" if is_db_ready() else "fallback",
            }
        except Exception as exc:
            import traceback
            traceback.print_exc()
            raise HTTPException(400, f"Failed to ingest CSV/TXT report: {exc}")

    # 2. Handle Excel Workbooks (.xlsx, .xlsm)
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not parse Excel workbook: {exc}")

    if is_db_ready():
        try:
            with db_connect(autocommit=True) as cx:
                loader = Loader(cx, wb, period_label, period_start, period_end)
                counts = loader.run()
                cx.execute("""
                    INSERT INTO etl_run (source_file, file_modified, finished_at, row_counts, notes)
                    VALUES (%s, now(), now(), %s, %s)
                """, (
                    fname,
                    json.dumps(counts),
                    f"Uploaded by {uploaded_by}" + (("; " + "; ".join(loader.warnings)) if loader.warnings else ""),
                ))

            try:
                await broker.notify("workbook_reload", "etl_run", {"counts": counts, "period": period_label})
            except Exception:
                pass

            return {
                "status": "success",
                "message": f"Successfully ingested '{fname}' into PostgreSQL database for {period_label}.",
                "filename": fname,
                "file_type": "excel",
                "period": period_label,
                "period_range": f"{period_start} to {period_end}",
                "uploaded_by": uploaded_by,
                "counts": counts,
                "warnings": loader.warnings,
                "environment": "postgres",
            }
        except Exception as exc:
            import traceback
            tb = traceback.format_exc()
            print("ERROR IN LOADER:", tb)
            raise HTTPException(500, f"Database ingestion failed: {exc}")

    # Fallback mode (when running in serverless / offline without DB):
    counts = {}
    sheet_map = {
        "Live Booking": "booking",
        "Vehicle Status": "vehicle",
        "Enquiries": "lead",
        "Enquiry": "lead",
        "Test Drive": "test_drive",
        "Allotment": "allotment",
        "Reg Report": "registration",
    }
    for sname in wb.sheetnames:
        for prefix, table in sheet_map.items():
            if prefix.lower() in sname.lower():
                ws = wb[sname]
                filled_rows = sum(
                    1 for r in range(2, min(ws.max_row + 1, 5000))
                    if ws.cell(r, 1).value or ws.cell(r, 2).value
                )
                counts[table] = max(counts.get(table, 0), filled_rows)

    return {
        "status": "success",
        "message": f"Successfully parsed '{fname}' for {period_label} (Serverless/Preview mode).",
        "filename": fname,
        "period": period_label,
        "period_range": f"{period_start} to {period_end}",
        "uploaded_by": uploaded_by,
        "counts": counts,
        "warnings": [],
        "environment": "serverless",
    }


@router.get("/api/etl-history", tags=["ingestion"])
def etl_history(limit: int = Query(10, le=50)):
    """Return recent Excel workbook ingestion runs for audit and status."""
    if is_db_ready():
        try:
            return fetch_all("""
                SELECT run_id, source_file, file_modified, finished_at, row_counts, notes
                FROM etl_run
                ORDER BY finished_at DESC
                LIMIT %s
            """, (limit,))
        except Exception:
            pass
    return [
        {
            "run_id": 1,
            "source_file": "DSR August 2026.xlsx",
            "file_modified": datetime.now().isoformat(),
            "finished_at": datetime.now().isoformat(),
            "row_counts": {"vehicle": 107, "lead": 2154, "booking": 133, "allotment": 20, "registration": 21},
            "notes": "Initial seed load",
        }
    ]
